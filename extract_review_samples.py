"""
从生成报告的 03_排查明细 中提取指定 Category 的差异 case，输出人工审核样本表 (CSV)。

典型用途（默认 finv=Rent）：
- miss category：illion 识别为 (空)，而 finv 识别为 Rent（illion 漏识别）
- incorrect category：illion 识别为其他类别，而 finv 识别为 Rent（illion 识别错类别）

用法：
    python extract_review_samples.py
    python extract_review_samples.py --category Rent --output rent_review_samples.csv
    python extract_review_samples.py --input category_difference_report_v3.xlsx

说明：
- 03_排查明细 的表头在第 3 行（第 1 行标题、第 2 行说明），脚本自动定位；
- finv Category 按精确值匹配（--category 区分大小写）；
- balance 字段由 label_compare.py 生成报告时补充；若报告未含该列则输出空列并提示。
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path
from typing import Any

import openpyxl

DEFAULT_INPUT = Path(__file__).resolve().parent / "category_difference_report_v2.xlsx"
DEFAULT_SHEET = "03_排查明细"
DEFAULT_CATEGORY = "Rent"
EMPTY_LABEL = "(空)"

# 输出的审核样本列：issue_type 列在最前，其余按人工审核的阅读顺序排列。
# 名称映射：报告表头 -> 输出列名（illion = 外部基准，finv = bscat = 内部模型）。
OUTPUT_COLUMNS = [
    "issue_type",
    "application_id",
    "job_id",
    "bank_account_id",
    "transaction_id",
    "transaction_date",
    "text",
    "amount",
    "balance",
    "dr_cr",
    "illion_category",
    "illion_third_party",
    "bscat_category",
    "bscat_third_party",
]

HEADER_RENAME = {
    "illion Category": "illion_category",
    "finv Category": "bscat_category",
    "third_party": "illion_third_party",
    "counterparty": "bscat_third_party",
}


def find_header_row(ws: openpyxl.worksheet.worksheet.Worksheet) -> int | None:
    """定位表头行号：03_排查明细 表头在第 3 行，但按内容识别更稳妥。

    表头行需同时包含 "illion Category" 与 "finv Category"。
    """
    for row in ws.iter_rows(min_row=1, max_row=8):
        values = [str(cell.value) for cell in row if cell.value is not None]
        if any(v.strip() == "illion Category" for v in values) and any(
            v.strip() == "finv Category" for v in values
        ):
            return row[0].row
    return None


def load_detail(ws: openpyxl.worksheet.worksheet.Worksheet) -> tuple[list[dict[str, Any]], str | None]:
    """读取 03_排查明细 内容，返回（标准化列名后的行列表，缺失的 balance 提示）。"""
    header_row = find_header_row(ws)
    if header_row is None:
        raise ValueError(
            f"在 {DEFAULT_SHEET} 前 8 行未找到表头（需包含 'illion Category' 和 'finv Category'），"
            "请确认输入文件是 label_compare.py 生成的报告"
        )

    # 报告表头 -> 输出列名（不在映射里的字段保留原名，如 transaction_date、text）。
    header = next(
        ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True)
    )
    col_map: dict[int, str] = {}
    for idx, value in enumerate(header):
        if value is None:
            continue
        name = str(value).strip()
        col_map[idx] = HEADER_RENAME.get(name, name)

    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if all(v is None for v in row):
            continue
        record = {col_map[i]: v for i, v in enumerate(row) if i in col_map}
        rows.append(record)

    warning = None
    if "balance" not in col_map.values():
        warning = (
            "报告 03_排查明细 中不含 balance 列（需要用修改后的 label_compare.py 重新生成报告），"
            "输出中 balance 为空"
        )
    return rows, warning


def is_blank(value: Any) -> bool:
    """识别空值：报告内空分类统一填充为 (空)，这里也兜底处理 NaN/None。"""
    if value is None:
        return True
    if isinstance(value, float):
        return value != value  # NaN
    return str(value).strip() in {"", EMPTY_LABEL}


def extract_cases(
    rows: list[dict[str, Any]],
    category: str,
) -> list[dict[str, Any]]:
    """按两个问题模式筛选 case，并打上 问题类型 标记。"""
    cases: list[dict[str, Any]] = []
    for record in rows:
        bscat = str(record.get("bscat_category", "")).strip()
        if bscat != category:
            continue
        illion = record.get("illion_category", "")
        if is_blank(illion):
            issue = "miss category"
        elif str(illion).strip() != category:
            issue = "incorrect category"
        else:
            continue  # 两侧都是 Rent，无差异，不属于本需求
        case = {"issue_type": issue, **record}
        cases.append(case)
    return cases


def write_csv(cases: list[dict[str, Any]], output: Path) -> None:
    """写 CSV，utf-8-sig 编码保证 Excel 直接打开中文不乱码。"""
    with output.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=OUTPUT_COLUMNS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(cases)


def main() -> None:
    parser = argparse.ArgumentParser(description="从 03_排查明细 提取指定 Category 的人工审核样本")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="生成报告路径（默认: category_difference_report_v2.xlsx）")
    parser.add_argument("--sheet", default=DEFAULT_SHEET, help="Sheet 名（默认: 03_排查明细）")
    parser.add_argument("--category", default=DEFAULT_CATEGORY, help="要提取的 finv Category 值，精确匹配（默认: Rent）")
    parser.add_argument("--output", type=Path, default=None, help="输出 CSV 路径（默认: {category}_review_samples.csv）")
    args = parser.parse_args()

    if not args.input.exists():
        raise FileNotFoundError(f"输入文件不存在: {args.input}")
    output = args.output or Path(f"{args.category}_review_samples.csv")

    wb = openpyxl.load_workbook(args.input, read_only=True, data_only=True)
    if args.sheet not in wb.sheetnames:
        raise ValueError(f"文件 {args.input} 中不存在 Sheet '{args.sheet}'，实际有: {wb.sheetnames}")
    ws = wb[args.sheet]

    rows, warning = load_detail(ws)
    if warning:
        print(f"[警告] {warning}")

    cases = extract_cases(rows, args.category)
    write_csv(cases, output)

    from collections import Counter

    counts = Counter(case["issue_type"] for case in cases)
    print(f"共提取 {len(cases)} 条 case 到 {output}")
    for issue in ["miss category", "incorrect category"]:
        if issue in counts:
            print(f"  {issue} → {counts[issue]} 条")


if __name__ == "__main__":
    main()
